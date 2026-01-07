# pyright: reportPrivateUsage=false
# pyright: reportUnknownMemberType=false

"""Unit test suite for the docx.revision module."""

from __future__ import annotations

import datetime as dt
from typing import cast

import pytest

from docx.oxml.ns import qn
from docx.oxml.revision import CT_RunTrackChange, CT_TrackChange
from docx.revision import TrackedChange, TrackedDeletion, TrackedInsertion
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from .unitutil.cxml import element, xml
from .unitutil.mock import FixtureRequest, Mock, instance_mock, property_mock


class DescribeCT_TrackChange:
    """Unit-test suite for `docx.oxml.revision.CT_TrackChange`."""

    def it_provides_access_to_the_id_attribute(self):
        ins = cast(CT_TrackChange, element("w:ins{w:id=42,w:author=John}"))
        assert ins.id == 42

    def it_can_set_the_id_attribute(self):
        ins = cast(CT_TrackChange, element("w:ins{w:id=1,w:author=John}"))
        ins.id = 99
        assert ins.id == 99

    def it_provides_access_to_the_author_attribute(self):
        ins = cast(CT_TrackChange, element("w:ins{w:id=1,w:author=Jane Doe}"))
        assert ins.author == "Jane Doe"

    def it_can_set_the_author_attribute(self):
        ins = cast(CT_TrackChange, element("w:ins{w:id=1,w:author=John}"))
        ins.author = "Jane Doe"
        assert ins.author == "Jane Doe"

    def it_provides_access_to_the_date_attribute(self):
        ins = cast(
            CT_TrackChange,
            element("w:ins{w:id=1,w:author=John,w:date=2024-01-15T10:30:00Z}"),
        )
        assert ins.date == "2024-01-15T10:30:00Z"

    def it_returns_None_when_date_attribute_is_not_present(self):
        ins = cast(CT_TrackChange, element("w:ins{w:id=1,w:author=John}"))
        assert ins.date is None

    def it_provides_date_value_as_datetime(self):
        ins = cast(
            CT_TrackChange,
            element("w:ins{w:id=1,w:author=John,w:date=2024-01-15T10:30:00Z}"),
        )
        date_val = ins.date_value
        assert date_val is not None
        assert date_val.year == 2024
        assert date_val.month == 1
        assert date_val.day == 15
        assert date_val.hour == 10
        assert date_val.minute == 30

    def it_returns_None_for_date_value_when_date_not_set(self):
        ins = cast(CT_TrackChange, element("w:ins{w:id=1,w:author=John}"))
        assert ins.date_value is None

    def it_can_set_date_value_from_datetime(self):
        ins = cast(CT_TrackChange, element("w:ins{w:id=1,w:author=John}"))
        ins.date_value = dt.datetime(2024, 6, 15, 14, 30, 0, tzinfo=dt.timezone.utc)
        assert ins.date == "2024-06-15T14:30:00Z"

    def it_can_clear_date_value_by_setting_None(self):
        ins = cast(
            CT_TrackChange,
            element("w:ins{w:id=1,w:author=John,w:date=2024-01-15T10:30:00Z}"),
        )
        ins.date_value = None
        assert ins.date is None


class DescribeCT_RunTrackChange:
    """Unit-test suite for `docx.oxml.revision.CT_RunTrackChange`."""

    def it_provides_access_to_paragraph_elements(self):
        ins = cast(CT_RunTrackChange, element("w:ins{w:id=1,w:author=John}/(w:p,w:p)"))
        assert len(ins.p_lst) == 2

    def it_provides_access_to_run_elements(self):
        ins = cast(CT_RunTrackChange, element("w:ins{w:id=1,w:author=John}/(w:r,w:r,w:r)"))
        assert len(ins.r_lst) == 3

    def it_provides_inner_content_elements_for_block_level_content(self):
        ins = cast(CT_RunTrackChange, element("w:ins{w:id=1,w:author=John}/(w:p,w:tbl,w:p)"))
        elements = ins.inner_content_elements
        assert len(elements) == 3
        assert elements[0].tag == qn("w:p")
        assert elements[1].tag == qn("w:tbl")
        assert elements[2].tag == qn("w:p")

    def it_provides_run_content_elements_for_run_level_content(self):
        ins = cast(CT_RunTrackChange, element("w:ins{w:id=1,w:author=John}/(w:r,w:r)"))
        elements = ins.run_content_elements
        assert len(elements) == 2
        assert elements[0].tag == qn("w:r")
        assert elements[1].tag == qn("w:r")


class DescribeTrackedInsertion:
    """Unit-test suite for `docx.revision.TrackedInsertion`."""

    def it_provides_access_to_author(self, parent_: Mock):
        ins_elm = cast(CT_RunTrackChange, element("w:ins{w:id=1,w:author=Alice}"))
        tracked = TrackedInsertion(ins_elm, parent_)
        assert tracked.author == "Alice"

    def it_can_set_author(self, parent_: Mock):
        ins_elm = cast(CT_RunTrackChange, element("w:ins{w:id=1,w:author=Alice}"))
        tracked = TrackedInsertion(ins_elm, parent_)
        tracked.author = "Bob"
        assert tracked.author == "Bob"

    def it_provides_access_to_revision_id(self, parent_: Mock):
        ins_elm = cast(CT_RunTrackChange, element("w:ins{w:id=42,w:author=Alice}"))
        tracked = TrackedInsertion(ins_elm, parent_)
        assert tracked.revision_id == 42

    def it_can_set_revision_id(self, parent_: Mock):
        ins_elm = cast(CT_RunTrackChange, element("w:ins{w:id=1,w:author=Alice}"))
        tracked = TrackedInsertion(ins_elm, parent_)
        tracked.revision_id = 99
        assert tracked.revision_id == 99

    def it_provides_access_to_date(self, parent_: Mock):
        ins_elm = cast(
            CT_RunTrackChange,
            element("w:ins{w:id=1,w:author=Alice,w:date=2024-03-15T08:00:00Z}"),
        )
        tracked = TrackedInsertion(ins_elm, parent_)
        date_val = tracked.date
        assert date_val is not None
        assert date_val.year == 2024
        assert date_val.month == 3

    def it_detects_block_level_content(self, parent_: Mock):
        ins_elm = cast(CT_RunTrackChange, element("w:ins{w:id=1,w:author=Alice}/w:p"))
        tracked = TrackedInsertion(ins_elm, parent_)
        assert tracked.is_block_level is True
        assert tracked.is_run_level is False

    def it_detects_run_level_content(self, parent_: Mock):
        ins_elm = cast(CT_RunTrackChange, element("w:ins{w:id=1,w:author=Alice}/w:r"))
        tracked = TrackedInsertion(ins_elm, parent_)
        assert tracked.is_block_level is False
        assert tracked.is_run_level is True

    def it_provides_access_to_runs(self, parent_: Mock):
        ins_elm = cast(
            CT_RunTrackChange,
            element("w:ins{w:id=1,w:author=Alice}/(w:r/w:t{xml:space=preserve}\"Hello\",w:r/w:t{xml:space=preserve}\" World\")"),
        )
        tracked = TrackedInsertion(ins_elm, parent_)
        runs = tracked.runs
        assert len(runs) == 2
        assert all(isinstance(r, Run) for r in runs)

    def it_provides_text_for_run_level_insertion(self, parent_: Mock):
        ins_elm = cast(
            CT_RunTrackChange,
            element("w:ins{w:id=1,w:author=Alice}/(w:r/w:t{xml:space=preserve}\"Hello\",w:r/w:t{xml:space=preserve}\" World\")"),
        )
        tracked = TrackedInsertion(ins_elm, parent_)
        assert tracked.text == "Hello World"

    def it_can_accept_insertion(self, parent_: Mock):
        body = element("w:body/(w:p,w:ins{w:id=1,w:author=Alice}/w:r/w:t\"inserted\",w:p)")
        ins_elm = cast(CT_RunTrackChange, body[1])
        tracked = TrackedInsertion(ins_elm, parent_)

        tracked.accept()

        assert body.xml == xml("w:body/(w:p,w:r/w:t\"inserted\",w:p)")

    def it_can_reject_insertion(self, parent_: Mock):
        body = element("w:body/(w:p,w:ins{w:id=1,w:author=Alice}/w:r/w:t\"inserted\",w:p)")
        ins_elm = cast(CT_RunTrackChange, body[1])
        tracked = TrackedInsertion(ins_elm, parent_)

        tracked.reject()

        assert body.xml == xml("w:body/(w:p,w:p)")

    @pytest.fixture
    def parent_(self, request: FixtureRequest):
        return instance_mock(request, Paragraph)


class DescribeTrackedDeletion:
    """Unit-test suite for `docx.revision.TrackedDeletion`."""

    def it_provides_access_to_author(self, parent_: Mock):
        del_elm = cast(CT_RunTrackChange, element("w:del{w:id=1,w:author=Bob}"))
        tracked = TrackedDeletion(del_elm, parent_)
        assert tracked.author == "Bob"

    def it_provides_access_to_revision_id(self, parent_: Mock):
        del_elm = cast(CT_RunTrackChange, element("w:del{w:id=55,w:author=Bob}"))
        tracked = TrackedDeletion(del_elm, parent_)
        assert tracked.revision_id == 55

    def it_detects_run_level_content(self, parent_: Mock):
        del_elm = cast(CT_RunTrackChange, element("w:del{w:id=1,w:author=Bob}/w:r"))
        tracked = TrackedDeletion(del_elm, parent_)
        assert tracked.is_run_level is True
        assert tracked.is_block_level is False

    def it_provides_text_for_run_level_deletion(self, parent_: Mock):
        del_elm = cast(
            CT_RunTrackChange,
            element("w:del{w:id=1,w:author=Bob}/w:r/w:t\"deleted text\""),
        )
        tracked = TrackedDeletion(del_elm, parent_)
        assert tracked.text == "deleted text"

    def it_can_accept_deletion(self, parent_: Mock):
        body = element("w:body/(w:p,w:del{w:id=1,w:author=Bob}/w:r/w:t\"deleted\",w:p)")
        del_elm = cast(CT_RunTrackChange, body[1])
        tracked = TrackedDeletion(del_elm, parent_)

        tracked.accept()

        assert body.xml == xml("w:body/(w:p,w:p)")

    def it_can_reject_deletion(self, parent_: Mock):
        body = element("w:body/(w:p,w:del{w:id=1,w:author=Bob}/w:r/w:t\"deleted\",w:p)")
        del_elm = cast(CT_RunTrackChange, body[1])
        tracked = TrackedDeletion(del_elm, parent_)

        tracked.reject()

        assert body.xml == xml("w:body/(w:p,w:r/w:t\"deleted\",w:p)")

    @pytest.fixture
    def parent_(self, request: FixtureRequest):
        return instance_mock(request, Paragraph)


class DescribeParagraph_iter_inner_content_with_revisions:
    """Unit-test suite for `Paragraph.iter_inner_content` with include_revisions=True."""

    def it_yields_tracked_insertions_when_include_revisions_is_True(
        self, document_part_: Mock
    ):
        p = element("w:p/(w:r/w:t\"normal\",w:ins{w:id=1,w:author=Alice}/w:r/w:t\"inserted\")")
        paragraph = Paragraph(p, document_part_)

        items = list(paragraph.iter_inner_content(include_revisions=True))

        assert len(items) == 2
        assert isinstance(items[0], Run)
        assert isinstance(items[1], TrackedInsertion)

    def it_yields_tracked_deletions_when_include_revisions_is_True(
        self, document_part_: Mock
    ):
        p = element("w:p/(w:r/w:t\"normal\",w:del{w:id=1,w:author=Bob}/w:r/w:t\"deleted\")")
        paragraph = Paragraph(p, document_part_)

        items = list(paragraph.iter_inner_content(include_revisions=True))

        assert len(items) == 2
        assert isinstance(items[0], Run)
        assert isinstance(items[1], TrackedDeletion)

    def it_excludes_revisions_by_default(self, document_part_: Mock):
        p = element("w:p/(w:r/w:t\"normal\",w:ins{w:id=1,w:author=Alice}/w:r/w:t\"inserted\")")
        paragraph = Paragraph(p, document_part_)

        items = list(paragraph.iter_inner_content())

        assert len(items) == 1
        assert isinstance(items[0], Run)

    @pytest.fixture
    def document_part_(self, request: FixtureRequest):
        from docx.parts.document import DocumentPart

        return instance_mock(request, DocumentPart)


class DescribeParagraph_add_run_tracked:
    """Unit-test suite for `Paragraph.add_run_tracked`."""

    def it_adds_a_tracked_insertion_with_text(self, document_part_: Mock):
        p = element("w:p")
        paragraph = Paragraph(p, document_part_)

        tracked = paragraph.add_run_tracked(text="new text", author="TestAuthor", revision_id=1)

        assert isinstance(tracked, TrackedInsertion)
        assert tracked.author == "TestAuthor"
        assert tracked.revision_id == 1
        ins_elements = p.xpath("./w:ins")
        assert len(ins_elements) == 1

    def it_auto_generates_revision_id_when_not_provided(self, document_part_: Mock):
        p = element("w:p/w:ins{w:id=5,w:author=Other}/w:r")
        paragraph = Paragraph(p, document_part_)

        tracked = paragraph.add_run_tracked(text="new", author="TestAuthor")

        assert tracked.revision_id == 6

    @pytest.fixture
    def document_part_(self, request: FixtureRequest):
        from docx.parts.document import DocumentPart

        return instance_mock(request, DocumentPart)


class DescribeRun_delete_tracked:
    """Unit-test suite for `Run.delete_tracked`."""

    def it_wraps_run_in_del_element(self, document_part_: Mock):
        p = element("w:p/w:r/w:t\"to delete\"")
        r = p[0]
        run = Run(r, document_part_)

        tracked = run.delete_tracked(author="Deleter", revision_id=10)

        assert isinstance(tracked, TrackedDeletion)
        assert tracked.author == "Deleter"
        assert tracked.revision_id == 10
        del_elements = p.xpath("./w:del")
        assert len(del_elements) == 1
        assert del_elements[0][0].tag == qn("w:r")

    def it_converts_w_t_to_w_delText(self, document_part_: Mock):
        p = element("w:p/w:r/w:t\"deleted text\"")
        r = p[0]
        run = Run(r, document_part_)

        run.delete_tracked(author="Deleter", revision_id=1)

        del_elem = p.xpath("./w:del")[0]
        r_elem = del_elem[0]
        t_elements = r_elem.xpath("./w:t")
        delText_elements = r_elem.xpath("./w:delText")
        assert len(t_elements) == 0
        assert len(delText_elements) == 1
        assert delText_elements[0].text == "deleted text"

    def it_auto_generates_revision_id_when_not_provided(self, document_part_: Mock):
        body = element("w:body/(w:del{w:id=7,w:author=Other}/w:r,w:p/w:r/w:t\"text\")")
        p = body[1]
        r = p[0]
        run = Run(r, document_part_)

        tracked = run.delete_tracked(author="Deleter")

        assert tracked.revision_id == 8

    @pytest.fixture
    def document_part_(self, request: FixtureRequest):
        from docx.parts.document import DocumentPart

        return instance_mock(request, DocumentPart)


class DescribeRun_replace_tracked_at:
    """Unit-test suite for `Run.replace_tracked_at`."""

    def it_replaces_text_at_specified_offsets(self, document_part_: Mock):
        p = element("w:p/w:r/w:t\"Hello World\"")
        r = p[0]
        run = Run(r, document_part_)

        run.replace_tracked_at(start=6, end=11, replace_text="Universe", author="Tester")

        del_elements = p.xpath(".//w:del")
        ins_elements = p.xpath(".//w:ins")
        assert len(del_elements) == 1
        assert len(ins_elements) == 1
        del_text = del_elements[0].xpath(".//w:delText")[0]
        ins_text = ins_elements[0].xpath(".//w:t")[0]
        assert del_text.text == "World"
        assert ins_text.text == "Universe"

    def it_preserves_text_before_replacement(self, document_part_: Mock):
        p = element("w:p/w:r/w:t\"Hello World\"")
        r = p[0]
        run = Run(r, document_part_)

        run.replace_tracked_at(start=6, end=11, replace_text="Universe", author="Tester")

        before_runs = p.xpath("./w:r/w:t[text()='Hello ']")
        assert len(before_runs) == 1

    def it_preserves_text_after_replacement(self, document_part_: Mock):
        p = element("w:p/w:r/w:t\"Hello World!\"")
        r = p[0]
        run = Run(r, document_part_)

        run.replace_tracked_at(start=6, end=11, replace_text="Universe", author="Tester")

        after_runs = p.xpath("./w:r/w:t[text()='!']")
        assert len(after_runs) == 1

    def it_handles_replacement_at_start_of_run(self, document_part_: Mock):
        p = element("w:p/w:r/w:t\"Hello World\"")
        r = p[0]
        run = Run(r, document_part_)

        run.replace_tracked_at(start=0, end=5, replace_text="Hi", author="Tester")

        del_text = p.xpath(".//w:delText")[0]
        ins_text = p.xpath(".//w:ins//w:t")[0]
        assert del_text.text == "Hello"
        assert ins_text.text == "Hi"

    def it_handles_replacement_at_end_of_run(self, document_part_: Mock):
        p = element("w:p/w:r/w:t\"Hello World\"")
        r = p[0]
        run = Run(r, document_part_)

        run.replace_tracked_at(start=6, end=11, replace_text="Everyone", author="Tester")

        del_text = p.xpath(".//w:delText")[0]
        ins_text = p.xpath(".//w:ins//w:t")[0]
        assert del_text.text == "World"
        assert ins_text.text == "Everyone"

    def it_raises_on_invalid_offsets(self, document_part_: Mock):
        p = element("w:p/w:r/w:t\"Hello\"")
        r = p[0]
        run = Run(r, document_part_)

        with pytest.raises(ValueError, match="Invalid offsets"):
            run.replace_tracked_at(start=10, end=15, replace_text="test", author="Tester")

    def it_raises_when_start_equals_end(self, document_part_: Mock):
        p = element("w:p/w:r/w:t\"Hello\"")
        r = p[0]
        run = Run(r, document_part_)

        with pytest.raises(ValueError, match="Invalid offsets"):
            run.replace_tracked_at(start=3, end=3, replace_text="test", author="Tester")

    def it_raises_when_start_greater_than_end(self, document_part_: Mock):
        p = element("w:p/w:r/w:t\"Hello\"")
        r = p[0]
        run = Run(r, document_part_)

        with pytest.raises(ValueError, match="Invalid offsets"):
            run.replace_tracked_at(start=4, end=2, replace_text="test", author="Tester")

    @pytest.fixture
    def document_part_(self, request: FixtureRequest):
        from docx.parts.document import DocumentPart

        return instance_mock(request, DocumentPart)


class DescribeParagraph_replace_tracked_at:
    """Unit-test suite for `Paragraph.replace_tracked_at`."""

    def it_replaces_text_within_single_run(self, document_part_: Mock):
        p = element("w:p/w:r/w:t\"Hello World\"")
        paragraph = Paragraph(p, document_part_)

        paragraph.replace_tracked_at(start=6, end=11, replace_text="Universe", author="Tester")

        del_elements = p.xpath(".//w:del")
        ins_elements = p.xpath(".//w:ins")
        assert len(del_elements) == 1
        assert len(ins_elements) == 1
        del_text = del_elements[0].xpath(".//w:delText")[0]
        ins_text = ins_elements[0].xpath(".//w:t")[0]
        assert del_text.text == "World"
        assert ins_text.text == "Universe"

    def it_replaces_text_spanning_multiple_runs(self, document_part_: Mock):
        p = element("w:p/(w:r/w:t\"Hello \",w:r/w:t\"World\")")
        paragraph = Paragraph(p, document_part_)

        paragraph.replace_tracked_at(start=4, end=9, replace_text="X", author="Tester")

        del_elements = p.xpath(".//w:del")
        ins_elements = p.xpath(".//w:ins")
        assert len(del_elements) == 1
        assert len(ins_elements) == 1
        del_text = del_elements[0].xpath(".//w:delText")[0]
        ins_text = ins_elements[0].xpath(".//w:t")[0]
        assert del_text.text == "o Wor"
        assert ins_text.text == "X"

    def it_preserves_text_before_and_after_multi_run_replacement(self, document_part_: Mock):
        p = element("w:p/(w:r/w:t\"Hello \",w:r/w:t\"World\")")
        paragraph = Paragraph(p, document_part_)

        paragraph.replace_tracked_at(start=4, end=9, replace_text="X", author="Tester")

        before_runs = p.xpath("./w:r/w:t[text()='Hell']")
        after_runs = p.xpath("./w:r/w:t[text()='ld']")
        assert len(before_runs) == 1
        assert len(after_runs) == 1

    def it_raises_on_invalid_offsets(self, document_part_: Mock):
        p = element("w:p/w:r/w:t\"Hello\"")
        paragraph = Paragraph(p, document_part_)

        with pytest.raises(ValueError, match="Invalid offsets"):
            paragraph.replace_tracked_at(start=10, end=15, replace_text="test", author="Tester")

    def it_raises_on_empty_paragraph(self, document_part_: Mock):
        p = element("w:p")
        paragraph = Paragraph(p, document_part_)

        with pytest.raises(ValueError, match="Invalid offsets"):
            paragraph.replace_tracked_at(start=0, end=5, replace_text="test", author="Tester")

    @pytest.fixture
    def document_part_(self, request: FixtureRequest):
        from docx.parts.document import DocumentPart

        return instance_mock(request, DocumentPart)


class DescribeParagraph_replace_tracked:
    """Unit-test suite for `Paragraph.replace_tracked`."""

    def it_replaces_text_at_word_level(self, document_part_: Mock):
        p = element("w:p/w:r/w:t\"Hello Unisys World\"")
        paragraph = Paragraph(p, document_part_)

        count = paragraph.replace_tracked("Unisys", "test", author="Tester")

        assert count == 1
        del_elements = p.xpath(".//w:del")
        ins_elements = p.xpath(".//w:ins")
        assert len(del_elements) == 1
        assert len(ins_elements) == 1
        del_text = del_elements[0].xpath(".//w:delText")[0]
        ins_text = ins_elements[0].xpath(".//w:t")[0]
        assert del_text.text == "Unisys"
        assert ins_text.text == "test"

    def it_handles_multiple_occurrences_in_same_run(self, document_part_: Mock):
        p = element("w:p/w:r/w:t\"Unisys and Unisys again\"")
        paragraph = Paragraph(p, document_part_)

        count = paragraph.replace_tracked("Unisys", "test", author="Tester")

        assert count == 2
        del_elements = p.xpath(".//w:del")
        ins_elements = p.xpath(".//w:ins")
        assert len(del_elements) == 2
        assert len(ins_elements) == 2

    def it_preserves_surrounding_text(self, document_part_: Mock):
        p = element("w:p/w:r/w:t\"Before Unisys After\"")
        paragraph = Paragraph(p, document_part_)

        paragraph.replace_tracked("Unisys", "test", author="Tester")

        all_text = "".join(t.text or "" for t in p.xpath(".//w:t | .//w:delText"))
        assert "Before" in all_text
        assert "After" in all_text
        assert "Unisys" in all_text
        assert "test" in all_text

    def it_returns_zero_when_no_match(self, document_part_: Mock):
        p = element("w:p/w:r/w:t\"No match here\"")
        paragraph = Paragraph(p, document_part_)

        count = paragraph.replace_tracked("Unisys", "test", author="Tester")

        assert count == 0

    @pytest.fixture
    def document_part_(self, request: FixtureRequest):
        from docx.parts.document import DocumentPart

        return instance_mock(request, DocumentPart)
